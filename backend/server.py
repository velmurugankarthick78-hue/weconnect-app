from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, status
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import csv
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional
import uuid
from datetime import datetime, timedelta, timezone
import bcrypt
import jwt as pyjwt
from openpyxl import load_workbook


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ.get("JWT_SECRET", "we-connect-dev-secret-change-in-prod")
JWT_ALGO = "HS256"
JWT_EXPIRES_HOURS = 24

ADMIN_EMAIL = "admin@weconnect.app"
ADMIN_PASSWORD = "WeConnect@2026"

app = FastAPI()
api_router = APIRouter(prefix="/api")
admin_router = APIRouter(prefix="/api/admin")
bearer = HTTPBearer(auto_error=False)


# ===== Models =====

class BusinessListing(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    businessName: str
    category: str
    ownerName: str
    mobileNumber: str
    whatsappNumber: Optional[str] = None
    address: str
    city: Optional[str] = ""
    district: str  # "Tirunelveli" | "Thoothukudi"
    googleMapsLink: Optional[str] = None
    imageUrl: Optional[str] = None
    description: Optional[str] = ""
    rating: Optional[float] = 4.5
    createdAt: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())


class BusinessCreate(BaseModel):
    businessName: str
    category: str
    ownerName: str
    mobileNumber: str
    whatsappNumber: Optional[str] = None
    address: str
    city: Optional[str] = ""
    district: str
    googleMapsLink: Optional[str] = None
    imageUrl: Optional[str] = None
    description: Optional[str] = ""


class BusinessUpdate(BaseModel):
    businessName: Optional[str] = None
    category: Optional[str] = None
    ownerName: Optional[str] = None
    mobileNumber: Optional[str] = None
    whatsappNumber: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    district: Optional[str] = None
    googleMapsLink: Optional[str] = None
    imageUrl: Optional[str] = None
    description: Optional[str] = None


class Category(BaseModel):
    key: str
    label: str
    icon: str
    imageUrl: str


class AdminLogin(BaseModel):
    email: str
    password: str


class TokenResponse(BaseModel):
    token: str
    email: str


class ImportResult(BaseModel):
    inserted: int
    skipped: int
    errors: List[str]


CATEGORIES: List[Category] = [
    Category(key="hotels", label="Hotels", icon="bed-outline",
             imageUrl="https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"),
    Category(key="hospitals", label="Hospitals", icon="medkit-outline",
             imageUrl="https://images.pexels.com/photos/668300/pexels-photo-668300.jpeg"),
    Category(key="battery", label="Battery Shops", icon="battery-charging-outline",
             imageUrl="https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"),
    Category(key="electricians", label="Electricians", icon="flash-outline",
             imageUrl="https://images.pexels.com/photos/9679179/pexels-photo-9679179.jpeg"),
    Category(key="plumbers", label="Plumbers", icon="water-outline",
             imageUrl="https://images.pexels.com/photos/6419128/pexels-photo-6419128.jpeg"),
    Category(key="auto", label="Auto Services", icon="car-outline",
             imageUrl="https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"),
    Category(key="restaurants", label="Restaurants", icon="restaurant-outline",
             imageUrl="https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"),
    Category(key="salons", label="Salons", icon="cut-outline",
             imageUrl="https://images.pexels.com/photos/3992874/pexels-photo-3992874.jpeg"),
    Category(key="grocery", label="Grocery", icon="basket-outline",
             imageUrl="https://images.pexels.com/photos/264636/pexels-photo-264636.jpeg"),
]


SEED_BUSINESSES = [
    # Tirunelveli
    {"businessName": "Iruttu Kadai Halwa Shop", "ownerName": "Selvaraj", "mobileNumber": "9442123456",
     "address": "West Car Street", "city": "Tirunelveli Town", "district": "Tirunelveli", "category": "restaurants",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7139,77.7567",
     "description": "Famous halwa shop since 1900s.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "Hotel Aryaas", "ownerName": "Murugan", "mobileNumber": "9842234567",
     "address": "67 Madurai Road", "city": "Tirunelveli Junction", "district": "Tirunelveli", "category": "hotels",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7256,77.6900",
     "description": "AC rooms with restaurant.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "Galaxy Hospital", "ownerName": "Dr. Karthik", "mobileNumber": "9842345678",
     "address": "Palayamkottai Main Road", "city": "Palayamkottai", "district": "Tirunelveli", "category": "hospitals",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7283,77.7341",
     "description": "Multi-speciality hospital, 24x7 emergency.",
     "imageUrl": "https://images.pexels.com/photos/668300/pexels-photo-668300.jpeg"},
    {"businessName": "Power Battery World", "ownerName": "Ramesh", "mobileNumber": "9943456789",
     "address": "Bypass Road, Vannarpettai", "city": "Tirunelveli", "district": "Tirunelveli", "category": "battery",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7194,77.7350",
     "description": "All brands - Exide, Amaron, Luminous.",
     "imageUrl": "https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"},
    {"businessName": "Sakthi Electricals", "ownerName": "Vinoth Kumar", "mobileNumber": "9789567890",
     "address": "Tower Junction", "city": "Tirunelveli Town", "district": "Tirunelveli", "category": "electricians",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7150,77.7480",
     "description": "Wiring, repairs, installations.",
     "imageUrl": "https://images.pexels.com/photos/9679179/pexels-photo-9679179.jpeg"},
    {"businessName": "Senthil Plumbing Works", "ownerName": "Senthil Kumar", "mobileNumber": "9787678901",
     "address": "VOC Nagar", "city": "Palayamkottai", "district": "Tirunelveli", "category": "plumbers",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7300,77.7350",
     "description": "24/7 emergency plumbing.",
     "imageUrl": "https://images.pexels.com/photos/6419128/pexels-photo-6419128.jpeg"},
    {"businessName": "Sri Lakshmi Auto Service", "ownerName": "Babu", "mobileNumber": "9043789012",
     "address": "NH-7 Bypass", "city": "Tirunelveli", "district": "Tirunelveli", "category": "auto",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7100,77.7100",
     "description": "Cars, bikes, all repairs.",
     "imageUrl": "https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"},
    {"businessName": "Aalayam Restaurant", "ownerName": "Mohan", "mobileNumber": "9842890123",
     "address": "Trivandrum Road", "city": "Tirunelveli", "district": "Tirunelveli", "category": "restaurants",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7220,77.7000",
     "description": "South Indian veg meals.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "Green Leaf Hospital", "ownerName": "Dr. Priya", "mobileNumber": "9786901234",
     "address": "Vannarpettai", "city": "Tirunelveli", "district": "Tirunelveli", "category": "hospitals",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7180,77.7380",
     "description": "Children & maternity care.",
     "imageUrl": "https://images.pexels.com/photos/668300/pexels-photo-668300.jpeg"},
    {"businessName": "Rolex Salon", "ownerName": "Anand", "mobileNumber": "9043012345",
     "address": "Trivandrum Road", "city": "Tirunelveli Junction", "district": "Tirunelveli", "category": "salons",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7250,77.6920",
     "description": "Unisex salon, AC.",
     "imageUrl": "https://images.pexels.com/photos/3992874/pexels-photo-3992874.jpeg"},
    {"businessName": "Saravana Stores", "ownerName": "Saravanan", "mobileNumber": "9842123450",
     "address": "Melapalayam", "city": "Tirunelveli", "district": "Tirunelveli", "category": "grocery",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7000,77.7300",
     "description": "Daily groceries, home delivery.",
     "imageUrl": "https://images.pexels.com/photos/264636/pexels-photo-264636.jpeg"},
    {"businessName": "Amaron Battery Point", "ownerName": "Karthik", "mobileNumber": "9994234561",
     "address": "Junction Main Road", "city": "Tirunelveli", "district": "Tirunelveli", "category": "battery",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7230,77.6950",
     "description": "Inverter & car batteries.",
     "imageUrl": "https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"},

    # Thoothukudi
    {"businessName": "Hotel Sea View", "ownerName": "Antony", "mobileNumber": "9842345671",
     "address": "Beach Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "hotels",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7642,78.1348",
     "description": "Sea-facing rooms, AC.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "Pearl City Hospital", "ownerName": "Dr. Stephen", "mobileNumber": "9943567812",
     "address": "Palayamkottai Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "hospitals",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7700,78.1300",
     "description": "Multi-speciality, ICU.",
     "imageUrl": "https://images.pexels.com/photos/668300/pexels-photo-668300.jpeg"},
    {"businessName": "Coastal Battery Shop", "ownerName": "Mariappan", "mobileNumber": "9789678923",
     "address": "Harbour Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "battery",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7650,78.1400",
     "description": "Marine & auto batteries.",
     "imageUrl": "https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"},
    {"businessName": "Selva Electricals", "ownerName": "Selvam", "mobileNumber": "9787789034",
     "address": "Old Bus Stand", "city": "Thoothukudi", "district": "Thoothukudi", "category": "electricians",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7625,78.1310",
     "description": "Home & commercial wiring.",
     "imageUrl": "https://images.pexels.com/photos/9679179/pexels-photo-9679179.jpeg"},
    {"businessName": "Royal Plumbing", "ownerName": "John Peter", "mobileNumber": "9043890145",
     "address": "Millerpuram", "city": "Thoothukudi", "district": "Thoothukudi", "category": "plumbers",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7710,78.1260",
     "description": "RO, motors, fittings.",
     "imageUrl": "https://images.pexels.com/photos/6419128/pexels-photo-6419128.jpeg"},
    {"businessName": "Marina Auto Garage", "ownerName": "Stephen Raj", "mobileNumber": "9842901256",
     "address": "Port Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "auto",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7660,78.1380",
     "description": "Cars & trucks, all brands.",
     "imageUrl": "https://images.pexels.com/photos/4374843/pexels-photo-4374843.jpeg"},
    {"businessName": "Macrons Restaurant", "ownerName": "Velu", "mobileNumber": "9786012367",
     "address": "Beach Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "restaurants",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7635,78.1370",
     "description": "Seafood specialities.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "Style Studio Salon", "ownerName": "Manoj", "mobileNumber": "9442123478",
     "address": "George Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "salons",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7680,78.1290",
     "description": "Hair, beard, facial.",
     "imageUrl": "https://images.pexels.com/photos/3992874/pexels-photo-3992874.jpeg"},
    {"businessName": "Annai Grocery Mart", "ownerName": "Devi", "mobileNumber": "9994234589",
     "address": "Anna Nagar", "city": "Thoothukudi", "district": "Thoothukudi", "category": "grocery",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7720,78.1340",
     "description": "Wholesale & retail.",
     "imageUrl": "https://images.pexels.com/photos/264636/pexels-photo-264636.jpeg"},
    {"businessName": "Sun Hotels", "ownerName": "Bharath", "mobileNumber": "9842345690",
     "address": "Trichy Road", "city": "Thoothukudi", "district": "Thoothukudi", "category": "hotels",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7600,78.1250",
     "description": "Budget AC stay.",
     "imageUrl": "https://images.pexels.com/photos/34487094/pexels-photo-34487094.jpeg"},
    {"businessName": "City Care Hospital", "ownerName": "Dr. Vinodh", "mobileNumber": "9043456712",
     "address": "Bryant Nagar", "city": "Thoothukudi", "district": "Thoothukudi", "category": "hospitals",
     "googleMapsLink": "https://www.google.com/maps/search/?api=1&query=8.7690,78.1320",
     "description": "General medicine, surgery.",
     "imageUrl": "https://images.pexels.com/photos/668300/pexels-photo-668300.jpeg"},
]


# ===== Auth Helpers =====

def hash_password(plain: str) -> str:
    return bcrypt.hashpw(plain.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(email: str) -> str:
    payload = {
        "sub": email,
        "exp": datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRES_HOURS),
        "iat": datetime.now(timezone.utc),
    }
    return pyjwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


async def require_admin(creds: Optional[HTTPAuthorizationCredentials] = Depends(bearer)) -> str:
    if creds is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Missing token")
    try:
        payload = pyjwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGO])
        email = payload.get("sub")
        if not email:
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")
        admin = await db.admins.find_one({"email": email}, {"_id": 0})
        if not admin:
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Admin not found")
        return email
    except pyjwt.ExpiredSignatureError:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Token expired")
    except pyjwt.InvalidTokenError:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")


# ===== Public Routes =====

@api_router.get("/")
async def root():
    return {"message": "We Connect API"}


@api_router.get("/categories", response_model=List[Category])
async def get_categories():
    return CATEGORIES


@api_router.get("/businesses", response_model=List[BusinessListing])
async def list_businesses(
    district: Optional[str] = None,
    category: Optional[str] = None,
    search: Optional[str] = None,
):
    query: dict = {}
    if district:
        query["district"] = district
    if category:
        query["category"] = category
    if search:
        query["$or"] = [
            {"businessName": {"$regex": search, "$options": "i"}},
            {"ownerName": {"$regex": search, "$options": "i"}},
            {"category": {"$regex": search, "$options": "i"}},
            {"address": {"$regex": search, "$options": "i"}},
            {"city": {"$regex": search, "$options": "i"}},
        ]
    items = await db.business_listings.find(query, {"_id": 0}).to_list(1000)
    return [BusinessListing(**item) for item in items]


@api_router.get("/businesses/{business_id}", response_model=BusinessListing)
async def get_business(business_id: str):
    item = await db.business_listings.find_one({"id": business_id}, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Business not found")
    return BusinessListing(**item)


@api_router.post("/businesses", response_model=BusinessListing)
async def create_business_public(payload: BusinessCreate):
    """Public form submission - same as admin create but unauthenticated."""
    biz = BusinessListing(**payload.dict())
    await db.business_listings.insert_one(biz.dict())
    return biz


# ===== Admin Routes =====

@admin_router.post("/login", response_model=TokenResponse)
async def admin_login(payload: AdminLogin):
    admin = await db.admins.find_one({"email": payload.email.lower().strip()}, {"_id": 0})
    if not admin or not verify_password(payload.password, admin["password_hash"]):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid credentials")
    token = create_token(admin["email"])
    return TokenResponse(token=token, email=admin["email"])


@admin_router.get("/me")
async def admin_me(email: str = Depends(require_admin)):
    return {"email": email}


@admin_router.get("/businesses", response_model=List[BusinessListing])
async def admin_list_businesses(
    search: Optional[str] = None,
    email: str = Depends(require_admin),
):
    query: dict = {}
    if search:
        query["$or"] = [
            {"businessName": {"$regex": search, "$options": "i"}},
            {"ownerName": {"$regex": search, "$options": "i"}},
            {"mobileNumber": {"$regex": search, "$options": "i"}},
            {"category": {"$regex": search, "$options": "i"}},
            {"address": {"$regex": search, "$options": "i"}},
            {"city": {"$regex": search, "$options": "i"}},
            {"district": {"$regex": search, "$options": "i"}},
        ]
    items = await db.business_listings.find(query, {"_id": 0}).sort("createdAt", -1).to_list(2000)
    return [BusinessListing(**item) for item in items]


@admin_router.post("/businesses", response_model=BusinessListing)
async def admin_create_business(payload: BusinessCreate, email: str = Depends(require_admin)):
    biz = BusinessListing(**payload.dict())
    await db.business_listings.insert_one(biz.dict())
    return biz


@admin_router.put("/businesses/{business_id}", response_model=BusinessListing)
async def admin_update_business(
    business_id: str,
    payload: BusinessUpdate,
    email: str = Depends(require_admin),
):
    updates = {k: v for k, v in payload.dict().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="No fields to update")
    result = await db.business_listings.update_one({"id": business_id}, {"$set": updates})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Business not found")
    item = await db.business_listings.find_one({"id": business_id}, {"_id": 0})
    return BusinessListing(**item)


@admin_router.delete("/businesses/{business_id}")
async def admin_delete_business(business_id: str, email: str = Depends(require_admin)):
    result = await db.business_listings.delete_one({"id": business_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Business not found")
    return {"deleted": True, "id": business_id}


@admin_router.post("/import", response_model=ImportResult)
async def admin_import(
    file: UploadFile = File(...),
    email: str = Depends(require_admin),
):
    """
    Import businesses from CSV or Excel.
    Expected columns (case-sensitive, order does not matter):
      businessName, category, ownerName, mobileNumber, whatsappNumber,
      address, city, district, googleMapsLink, imageUrl
    """
    raw = await file.read()
    filename = (file.filename or "").lower()
    rows: List[dict] = []
    errors: List[str] = []

    try:
        if filename.endswith(".csv"):
            text = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                rows.append({(k or "").strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items()})
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            wb = load_workbook(io.BytesIO(raw), data_only=True)
            ws = wb.active
            if ws is None:
                raise ValueError("Empty workbook")
            iter_rows = ws.iter_rows(values_only=True)
            try:
                header_row = next(iter_rows)
            except StopIteration:
                header_row = ()
            headers = [str(h).strip() if h is not None else "" for h in header_row]
            for r in iter_rows:
                if r is None:
                    continue
                row = {}
                for i, h in enumerate(headers):
                    if not h:
                        continue
                    v = r[i] if i < len(r) else None
                    row[h] = str(v).strip() if v is not None else ""
                if any(row.values()):
                    rows.append(row)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file. Use .csv or .xlsx")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {e}")

    inserted = 0
    skipped = 0
    required = ["businessName", "category", "ownerName", "mobileNumber", "address", "district"]
    for idx, row in enumerate(rows, start=2):  # row 1 is header
        missing = [k for k in required if not str(row.get(k, "")).strip()]
        if missing:
            skipped += 1
            errors.append(f"Row {idx}: missing {', '.join(missing)}")
            continue
        try:
            biz = BusinessListing(
                businessName=str(row.get("businessName", "")).strip(),
                category=str(row.get("category", "")).strip(),
                ownerName=str(row.get("ownerName", "")).strip(),
                mobileNumber=str(row.get("mobileNumber", "")).strip(),
                whatsappNumber=str(row.get("whatsappNumber", "")).strip() or None,
                address=str(row.get("address", "")).strip(),
                city=str(row.get("city", "")).strip(),
                district=str(row.get("district", "")).strip(),
                googleMapsLink=str(row.get("googleMapsLink", "")).strip() or None,
                imageUrl=str(row.get("imageUrl", "")).strip() or None,
                description=str(row.get("description", "")).strip(),
            )
            await db.business_listings.insert_one(biz.dict())
            inserted += 1
        except Exception as e:
            skipped += 1
            errors.append(f"Row {idx}: {e}")

    return ImportResult(inserted=inserted, skipped=skipped, errors=errors[:50])


# ===== App Setup =====

app.include_router(api_router)
app.include_router(admin_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


@app.on_event("startup")
async def startup_seed():
    try:
        # Seed admin (idempotent)
        await db.admins.create_index("email", unique=True)
        existing = await db.admins.find_one({"email": ADMIN_EMAIL})
        if not existing:
            await db.admins.insert_one({
                "email": ADMIN_EMAIL,
                "password_hash": hash_password(ADMIN_PASSWORD),
                "createdAt": datetime.now(timezone.utc).isoformat(),
            })
            logger.info(f"Seeded admin {ADMIN_EMAIL}")

        # Migrate legacy `businesses` collection -> `business_listings`
        legacy_count = await db.businesses.count_documents({})
        new_count = await db.business_listings.count_documents({})
        if new_count == 0 and legacy_count > 0:
            async for old in db.businesses.find({}, {"_id": 0}):
                migrated = {
                    "id": old.get("id", str(uuid.uuid4())),
                    "businessName": old.get("name", ""),
                    "category": old.get("category", ""),
                    "ownerName": old.get("owner_name", ""),
                    "mobileNumber": old.get("mobile", ""),
                    "whatsappNumber": old.get("whatsapp") or None,
                    "address": old.get("address", ""),
                    "city": "",
                    "district": old.get("district", ""),
                    "googleMapsLink": (
                        f"https://www.google.com/maps/search/?api=1&query={old['latitude']},{old['longitude']}"
                        if old.get("latitude") and old.get("longitude") else None
                    ),
                    "imageUrl": old.get("image_url"),
                    "description": old.get("description", ""),
                    "rating": old.get("rating", 4.5),
                    "createdAt": old.get("created_at", datetime.now(timezone.utc).isoformat()),
                }
                await db.business_listings.insert_one(migrated)
            logger.info(f"Migrated {legacy_count} legacy businesses -> business_listings")
            new_count = await db.business_listings.count_documents({})

        # Seed business_listings if still empty
        if new_count == 0:
            docs = [BusinessListing(**s).dict() for s in SEED_BUSINESSES]
            if docs:
                await db.business_listings.insert_many(docs)
                logger.info(f"Seeded {len(docs)} business_listings")
    except Exception as e:
        logger.error(f"Seed/migrate failed: {e}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
